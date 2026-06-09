"""[검증] v1 자동 결과 (workspace/test_kim_21th_v1.xlsx) vs 김태웅 정답 (reference/21th_kim.xlsx).

검증 대상 시트: 전체, 100미만, 100-300, 300-500, 500-1000, 1000이상
검증 대상 컬럼 그룹: 2)인력규모, 3)예상매출, 4)교육예산

(1)업종은 김태웅이 수기 재분류한 결과라 일치 기대 안 함 — 별도 분석)

산출: tests/test_stage4_report.md
"""
from pathlib import Path
import sys, io
if __name__ == "__main__" and sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import pandas as pd
import openpyxl as op

ROOT = Path(__file__).resolve().parents[1]
V1 = ROOT / "workspace/test_kim_21th_v1.xlsx"
KIM = ROOT / "reference/21th_kim.xlsx"
REPORT = ROOT / "tests/test_stage4_report.md"

KIM_FREQ_DATA_START = 5
KIM_TOTAL_ROW = 24

# 김태웅 시트 col 위치 (0-indexed, openpyxl 은 1-indexed)
# col 6: 1)업종 라벨, col 7: 빈도
# col 8: 2)인력규모 라벨, col 9: 빈도
# col 10: 3)예상매출 라벨, col 11: 빈도
# col 12: 4)교육예산 라벨, col 13: 빈도
GROUPS = {
    "2)인력규모": (8, 9),
    "3)예상매출": (10, 11),
    "4)교육예산": (12, 13),
}
SHEETS_TO_CHECK = ["전체", "100미만", "100-300", "300-500", "500-1000", "1000이상"]


def read_kim_block(wb, sheet_name: str) -> dict:
    """김태웅 시트에서 컬럼 그룹별 (라벨 → 빈도) 추출."""
    ws = wb[sheet_name]
    out = {}
    for group_name, (lbl_col, val_col) in GROUPS.items():
        d = {}
        # row 5부터 합계 전까지
        for r in range(KIM_FREQ_DATA_START, KIM_TOTAL_ROW):
            lbl = ws.cell(row=r, column=lbl_col + 1).value
            val = ws.cell(row=r, column=val_col + 1).value
            if lbl is None or (isinstance(lbl, str) and not lbl.strip()):
                continue
            if val is None:
                val = 0
            try:
                val = int(float(val))
            except (TypeError, ValueError):
                val = 0
            d[str(lbl).strip()] = val
        # 합계
        total = ws.cell(row=KIM_TOTAL_ROW, column=val_col + 1).value
        try:
            total = int(float(total)) if total else sum(d.values())
        except (TypeError, ValueError):
            total = sum(d.values())
        out[group_name] = {"counts": d, "total": total}
    return out


def read_v1_block(wb, sheet_name: str) -> dict:
    """v1 자동 결과에서 같은 블록 추출 (동일 구조)."""
    return read_kim_block(wb, sheet_name)  # 같은 형식


def compare_blocks(v1: dict, kim: dict, sheet: str) -> tuple[list[str], dict]:
    lines = [f"### {sheet}"]
    lines.append("")
    v_N = sum(v1.get("2)인력규모", {}).get("counts", {}).values())
    k_N = sum(kim.get("2)인력규모", {}).get("counts", {}).values())
    lines.append(f"- v1 N = {v_N} / 김태웅 N = {k_N} / Δ = {v_N - k_N:+d}")
    lines.append("")

    stats = {}
    for group_name in GROUPS:
        v_data = v1.get(group_name, {}).get("counts", {})
        k_data = kim.get(group_name, {}).get("counts", {})
        all_keys = sorted(set(v_data.keys()) | set(k_data.keys()))
        lines.append(f"#### {group_name}")
        lines.append("")
        lines.append("| 라벨 | v1 빈도 | 김태웅 빈도 | Δ빈도 | v1 비율 | 김태웅 비율 | Δ비율 |")
        lines.append("|---|---:|---:|---:|---:|---:|---:|")

        freq_match = 0
        ratio_match_2pp = 0
        ratio_match_5pp = 0
        for k in all_keys:
            vf = v_data.get(k, 0)
            kf = k_data.get(k, 0)
            df_ = vf - kf
            vr = vf / v_N if v_N else 0
            kr = kf / k_N if k_N else 0
            dr = (vr - kr) * 100  # pp
            fm = "✓" if df_ == 0 else ("△" if abs(df_) <= 2 else "✗")
            rm = "✓" if abs(dr) <= 2 else ("△" if abs(dr) <= 5 else "✗")
            if df_ == 0:
                freq_match += 1
            if abs(dr) <= 2:
                ratio_match_2pp += 1
            if abs(dr) <= 5:
                ratio_match_5pp += 1
            lines.append(
                f"| {k} | {vf} | {kf} | {df_:+d} {fm} | {vr*100:.1f}% | {kr*100:.1f}% | {dr:+.1f}pp {rm} |"
            )
        n = len(all_keys)
        lines.append(
            f"| **일치율** | | | **{freq_match}/{n}** | | | **±2pp: {ratio_match_2pp}/{n} · ±5pp: {ratio_match_5pp}/{n}** |"
        )
        lines.append("")
        stats[group_name] = {
            "n": n,
            "freq_match": freq_match,
            "ratio_match_2pp": ratio_match_2pp,
            "ratio_match_5pp": ratio_match_5pp,
        }
    return lines, stats


def main():
    wb_v1 = op.load_workbook(V1, data_only=True)
    wb_kim = op.load_workbook(KIM, data_only=True)

    print(f"v1 시트: {wb_v1.sheetnames}")
    print(f"김태웅 시트: {wb_kim.sheetnames}")

    report = [
        "# 단계 4 v1 검증 리포트 (21회 reverse-run)",
        "",
        "**비교**: `workspace/test_kim_21th_v1.xlsx` (v1 자동 결과) vs `reference/21th_kim.xlsx` (김태웅 정답)",
        "",
        "**검증 대상**: 6 시트 × 3 컬럼 그룹 (인력규모/예상매출/교육예산)",
        "",
        "**제외**: 1)업종 = 김태웅 수기 재분류 (자동 일치 기대 안 함). 5)예산변동 = raw 매핑 추가 검증 필요 (v1 미포함).",
        "",
    ]
    all_stats = {}
    for sheet in SHEETS_TO_CHECK:
        if sheet not in wb_v1.sheetnames or sheet not in wb_kim.sheetnames:
            report.append(f"### {sheet} — SKIP (시트 누락)")
            continue
        v1_block = read_v1_block(wb_v1, sheet)
        kim_block = read_kim_block(wb_kim, sheet)
        lines, stats = compare_blocks(v1_block, kim_block, sheet)
        report += lines
        all_stats[sheet] = stats

    # 종합 요약
    report.append("## 종합 요약 (비율 기반)")
    report.append("")
    report.append("> 빈도 기반은 모집단 크기 차이(v1 237 vs 김태웅 206 = +31 분석가 수기 분류 영향)로 차이 발생.")
    report.append("> **비율 기반**이 자동화 정확도 평가에 더 적합.")
    report.append("")
    report.append("| 시트 | 인력규모 ±2pp | 예상매출 ±2pp | 교육예산 ±2pp |")
    report.append("|---|---|---|---|")
    for sh, stats in all_stats.items():
        cells = []
        for grp in ["2)인력규모", "3)예상매출", "4)교육예산"]:
            s = stats.get(grp, {})
            cells.append(f"{s.get('ratio_match_2pp', 0)}/{s.get('n', 0)}")
        report.append(f"| {sh} | {' | '.join(cells)} |")
    report.append("")

    # ±5pp 종합
    report.append("| 시트 | 인력규모 ±5pp | 예상매출 ±5pp | 교육예산 ±5pp |")
    report.append("|---|---|---|---|")
    for sh, stats in all_stats.items():
        cells = []
        for grp in ["2)인력규모", "3)예상매출", "4)교육예산"]:
            s = stats.get(grp, {})
            cells.append(f"{s.get('ratio_match_5pp', 0)}/{s.get('n', 0)}")
        report.append(f"| {sh} | {' | '.join(cells)} |")
    report.append("")

    REPORT.write_text("\n".join(report), encoding="utf-8")
    print(f"\n리포트: {REPORT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
