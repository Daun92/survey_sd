"""[단계 4-세그] 범용 세그먼트 집계기 — 2026 설문 전 문항 × 임의 세그먼트 df.

config/segment_items_22th.yaml 의 문항 정의 + 코드북 0_survey/codebook_22th.xlsx 의
라벨·응답형식으로, 세그먼트 DataFrame 하나를 받아 전 문항(54) 집계값을 산출한다.
④ XLSX(stage4_build_kim_xlsx) · ⑤ HTML(stage5_render_segment_report) 공통 데이터 소스.

라벨은 config 에 두지 않고 코드북에서 런타임 도출:
  - dist/dist_multi/multi_rate : '코드표' 응답형식 'n=라벨' 파싱
  - mean/mean_pct/mean_pct_cf/likert/likert_iv/tenure_ym : '코드표' 하위항목
  - issue_rank/issue_multi : 'HRD이슈_항목코드'(A그룹 코드1~15 / B그룹 코드1~15)

신규 agg: mean_pct_cf · dist_multi · likert_iv · tenure_ym · issue_rank · issue_multi.

사용(단독 sanity):
    python scripts/segment_aggregate.py --src 2_validated/raw_validated.xlsx
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
from pathlib import Path

import openpyxl as op
import pandas as pd
import yaml

if __name__ == "__main__" and sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from stage4_extract_chart_data import num                       # noqa: E402
from stage5_render_html import freq_dist, multi_dist            # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
CODEBOOK = ROOT / "0_survey/codebook_22th.xlsx"
CONFIG = ROOT / "config/segment_items_22th.yaml"
DEFAULT_SRC = ROOT / "2_validated/raw_validated.xlsx"

# budget_chg(p1)R_2_n_2) — 응답형식이 비표준('1~5=증가(...)') → 명시(stage5 p6 과 동일 인코딩)
BUDGET_CHG_LABELS = {
    1: "증가 1~5%", 2: "증가 6~15%", 3: "증가 16~25%", 4: "증가 26~50%", 5: "증가 51%+",
    6: "감소 1~5%", 7: "감소 6~15%", 8: "감소 16~25%", 9: "감소 26~50%", 10: "감소 51%+",
    11: "동일",
}


# ─────────────────────────────────────────────
# 코드북
# ─────────────────────────────────────────────

def _s(v) -> str:
    return "" if v is None else str(v).strip()


class Codebook:
    """'코드표' + 'HRD이슈_항목코드' 인덱스."""

    def __init__(self, path: Path = CODEBOOK):
        wb = op.load_workbook(path, read_only=True)
        ws = wb["코드표"]
        self.rows: dict[str, dict] = {}
        self.order: list[str] = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue
            col = _s(r[0])
            if not col:
                continue
            self.rows[col] = {
                "part": _s(r[1]), "q": _s(r[2]), "sub": _s(r[3]),
                "fmt": _s(r[4]), "map21": _s(r[5]) if len(r) > 5 else "",
            }
            self.order.append(col)

        # HRD이슈 항목코드: (group 'A'/'B', code) -> 라벨
        self.issue: dict[tuple[str, int], str] = {}
        ws2 = wb["HRD이슈_항목코드"]
        for i, r in enumerate(ws2.iter_rows(values_only=True)):
            if i == 0:
                continue
            code, name, grp = r[0], r[1], r[2] if len(r) > 2 else None
            if code is None or name is None:
                continue
            g = "A" if "A그룹" in _s(grp) else ("B" if "B그룹" in _s(grp) else None)
            if g is None:
                continue
            try:
                ci = int(float(code))
            except (TypeError, ValueError):
                continue
            self.issue[(g, ci)] = _s(name)
        wb.close()

    def fmt(self, col: str) -> str:
        return self.rows.get(col, {}).get("fmt", "")

    def sub(self, col: str) -> str:
        return self.rows.get(col, {}).get("sub", "") or col

    def resolve_cols(self, cols_spec: list[str], df_cols: set) -> list[str]:
        """config cols → 실제 컬럼(코드북 순서). 와일드카드 전개 + 텍스트(_etc) 제외."""
        out: list[str] = []
        for c in cols_spec:
            if "*" in c:
                pre = c.replace("*", "")
                for x in self.order:
                    if x.startswith(pre) and x in df_cols and not self.fmt(x).startswith("텍스트"):
                        out.append(x)
            elif c in df_cols:
                out.append(c)
        return out


def parse_resp_labels(fmt: str) -> dict[int, str]:
    """'1=대기업, 2=중기업, ...' → {1:'대기업', ...}. 비표준 토큰은 무시(부분 파싱)."""
    out: dict[int, str] = {}
    for part in str(fmt).split(","):
        m = re.match(r"\s*(\d+)\s*=\s*(.+?)\s*$", part)
        if m:
            out[int(m.group(1))] = m.group(2).strip()
    return out


# ─────────────────────────────────────────────
# 라벨 정리 / 작은 헬퍼
# ─────────────────────────────────────────────

_CF_RE = re.compile(r"\s*-\s*(현재|향후)\s*비율\s*$")
_IV_RE = re.compile(r"\s*-\s*(중요도|수행도)\s*$")


def _r(v, nd: int = 2):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return round(float(v), nd)


def _clean_cf(s: str) -> str:
    return _CF_RE.sub("", str(s)).strip()


def _clean_iv(s: str) -> str:
    return _IV_RE.sub("", str(s)).strip()


def _score_dist(s) -> dict[int, float]:
    s = num(s)
    n = int(s.notna().sum())
    return {k: (_r(100 * int((s == k).sum()) / n, 1) if n else 0.0) for k in range(1, 6)}


def _parse_ym(v):
    """'년-월' 형식(예 '10-0','9-6') → 개월. 숫자만이면 년으로 간주."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    s = str(v).strip()
    m = re.match(r"^(\d+)\s*-\s*(\d+)$", s)
    if m:
        return int(m.group(1)) * 12 + int(m.group(2))
    try:
        return float(s) * 12
    except ValueError:
        return None


# ─────────────────────────────────────────────
# agg 구현
# ─────────────────────────────────────────────

def _agg_macro(df, cols, cb):
    items = []
    for c in cols:
        s = num(df[c])
        items.append({"col": c, "label": cb.sub(c), "mean": _r(s.mean()),
                      "median": _r(s.median()), "n": int(s.notna().sum())})
    return {"items": items}


def _agg_mean(df, cols, cb, unit="명"):
    items = []
    for c in cols:
        s = num(df[c])
        items.append({"col": c, "label": cb.sub(c), "mean": _r(s.mean()),
                      "median": _r(s.median()), "n": int(s.notna().sum())})
    return {"items": items, "unit": unit}


def _agg_tenure(df, cols, cb):
    items = []
    for c in cols:
        months = pd.to_numeric(df[c].map(_parse_ym), errors="coerce")
        mm, md = months.mean(), months.median()
        items.append({
            "col": c, "label": cb.sub(c), "n": int(months.notna().sum()),
            "mean_months": _r(mm, 1), "median_months": _r(md, 1),
            "mean_years": _r(mm / 12 if pd.notna(mm) else None, 1),
            "median_years": _r(md / 12 if pd.notna(md) else None, 1),
        })
    return {"items": items, "unit": "개월"}


def _pct_means(df, cols):
    """합~100 그룹 일관 분모(그룹합>0 행) 평균. (col->pct, base_N)."""
    if not cols:
        return {}, 0
    sub = df[cols].apply(num)
    base = sub[sub.sum(axis=1, min_count=1) > 0]
    if len(base) == 0:
        return {c: 0.0 for c in cols}, 0
    base = base.fillna(0)
    return {c: _r(base[c].mean()) for c in cols}, int(len(base))


def _agg_mean_pct(df, cols, cb):
    means, bn = _pct_means(df, cols)
    items = [{"col": c, "label": cb.sub(c), "pct": means.get(c, 0.0)} for c in cols]
    return {"items": items, "base_N": bn, "sum": _r(sum(x["pct"] or 0 for x in items), 1)}


def _agg_mean_pct_cf(df, cols, cb):
    cur = [c for c in cols if c.endswith("_3")]
    fut = [c for c in cols if c.endswith("_4")]
    cm, cn = _pct_means(df, cur)
    fm, fn = _pct_means(df, fut)
    keys = []
    for c in cur:
        k = re.sub(r"_3$", "", c)
        if k not in keys:
            keys.append(k)
    items = []
    for k in keys:
        c3, c4 = k + "_3", k + "_4"
        lab = _clean_cf(cb.sub(c3))
        items.append({"key": k, "label": lab, "cur": cm.get(c3), "fut": fm.get(c4),
                      "delta": _r((fm.get(c4) or 0) - (cm.get(c3) or 0))})
    items = _merge_dup_cf(items)
    return {"items": items, "base_N_cur": cn, "base_N_fut": fn}


def _merge_dup_cf(items):
    """동일 라벨(예 저성과자 교육 _7/_9) cur/fut 합산."""
    seen, out = {}, []
    for it in items:
        if it["label"] in seen:
            t = seen[it["label"]]
            for f in ("cur", "fut"):
                t[f] = _r((t.get(f) or 0) + (it.get(f) or 0))
            t["delta"] = _r((t.get("fut") or 0) - (t.get("cur") or 0))
            t["merged"] = t.get("merged", 1) + 1
        else:
            seen[it["label"]] = it
            out.append(it)
    return out


def _agg_dist(df, cols, cb, q):
    c = cols[0]
    lm = BUDGET_CHG_LABELS if q["id"] == "budget_chg" else parse_resp_labels(cb.fmt(c))
    items = freq_dist(df[c], lm)
    if q["id"] == "budget_chg":          # NaN(무응답) → 동일 로 흡수 (stage5 p6 정합)
        items = _fold_label(items, "무응답", "동일")
    return {"items": items, "col": c}


def _fold_label(items, src, dst):
    out, acc = [], None
    for it in items:
        if it["label"] == src:
            acc = it
            continue
        out.append(it)
    if acc:
        tgt = next((x for x in out if x["label"] == dst), None)
        if tgt:
            tgt["count"] += acc["count"]
            tgt["pct"] = _r(tgt["pct"] + acc["pct"], 1)
        else:
            acc["label"] = dst
            out.append(acc)
    return out


def _agg_dist_multi(df, cols, cb):
    items = []
    for c in cols:
        lm = parse_resp_labels(cb.fmt(c))
        dist = freq_dist(df[c], lm)
        impl = next((x["pct"] for x in dist if x["label"] == lm.get(4)), 0.0)  # 4=실시중
        items.append({"col": c, "label": cb.sub(c), "dist": dist, "impl_rate": impl})
    return {"items": items}


def _agg_multi_rate(df, cols, cb):
    c = cols[0]
    md = multi_dist(df[c], parse_resp_labels(cb.fmt(c)))
    return {"items": md["items"], "n_respondents": md["n_respondents"], "col": c}


def _agg_likert(df, cols, cb):
    items = []
    for c in cols:
        s = num(df[c])
        items.append({"col": c, "label": cb.sub(c), "mean": _r(s.mean()),
                      "n": int(s.notna().sum()), "dist": _score_dist(s)})
    return {"items": items}


def _agg_likert_iv(df, cols, cb):
    # 페어 모드: 모든 col 이 _3/_4 로 끝나고 item key 별 _3·_4 둘 다 존재 (중요도/수행도)
    keys: dict[str, dict] = {}
    paired = bool(cols)
    for c in cols:
        m = re.match(r"^(.*)_([34])$", c)
        if not m:
            paired = False
            break
        keys.setdefault(m.group(1), {})[m.group(2)] = c
    if paired and all("3" in v and "4" in v for v in keys.values()):
        items = []
        for k, v in keys.items():
            imp, perf = num(df[v["3"]]), num(df[v["4"]])
            items.append({"key": k, "label": _clean_iv(cb.sub(v["3"])),
                          "imp_mean": _r(imp.mean()), "perf_mean": _r(perf.mean()),
                          "gap": _r((imp.mean() or 0) - (perf.mean() or 0)),
                          "imp_dist": _score_dist(imp), "perf_dist": _score_dist(perf),
                          "n": int(imp.notna().sum())})
        return {"mode": "pairs", "items": items}
    # duo 모드: 2개 col 을 각각 한 항목으로 (중요성 vs 활용 / 만족 vs 활용)
    items = [{"col": c, "label": cb.sub(c), "mean": _r(num(df[c]).mean()),
              "dist": _score_dist(df[c]), "n": int(num(df[c]).notna().sum())} for c in cols]
    gap = (_r((num(df[cols[0]]).mean() or 0) - (num(df[cols[1]]).mean() or 0))
           if len(cols) >= 2 else None)
    return {"mode": "duo", "items": items, "gap": gap}


_GRP = {2: "A", 4: "B"}


def _agg_issue_rank(df, q, cb):
    """cheOrder='항목-그룹-순위위치' 콤마구분(그룹 2=A·4=B, 순위 1~3) 파싱."""
    order_col = next((c for c in q["cols"] if c.endswith("cheOrder")), None)
    counts: dict[tuple[str, int], dict[int, int]] = {}
    n_resp = 0
    if order_col and order_col in df.columns:
        for v in df[order_col].dropna():
            s = str(v).strip()
            if not s or s.lower() == "nan":
                continue
            n_resp += 1
            for tok in s.split(","):
                p = tok.strip().split("-")
                if len(p) != 3:
                    continue
                try:
                    item, grp, pos = int(p[0]), int(p[1]), int(p[2])
                except ValueError:
                    continue
                g = _GRP.get(grp)
                if g is None or pos not in (1, 2, 3):
                    continue
                counts.setdefault((g, item), {1: 0, 2: 0, 3: 0})[pos] += 1
    themes = []
    for (g, item), rc in counts.items():
        total = rc[1] + rc[2] + rc[3]
        themes.append({"group": g, "code": item, "label": cb.issue.get((g, item), f"{g}{item}"),
                       "r1": rc[1], "r2": rc[2], "r3": rc[3], "total": total,
                       "mention_pct": _r(100 * total / n_resp if n_resp else 0, 1),
                       "weighted": rc[1] * 3 + rc[2] * 2 + rc[3]})
    themes.sort(key=lambda x: -x["weighted"])
    return {"themes": themes, "n_respondents": n_resp}


def _agg_issue_multi(df, q, cb):
    a_col = next((c for c in q["cols"] if c.endswith("_n_2_che")), None)
    b_col = next((c for c in q["cols"] if c.endswith("_n_4_che")), None)
    mask = pd.Series(False, index=df.index)
    for c in (a_col, b_col):
        if c and c in df.columns:
            mask |= df[c].astype(str).str.strip().str.lower().replace("nan", "").ne("")
    n_resp = int(mask.sum())
    out = []

    def tally(col, g):
        if not col or col not in df.columns:
            return
        cnt: dict[int, int] = {}
        for v in df[col].dropna():
            for tok in str(v).split(","):
                tok = tok.strip()
                if not tok or tok.lower() == "nan":
                    continue
                try:
                    ci = int(float(tok))
                except ValueError:
                    continue
                cnt[ci] = cnt.get(ci, 0) + 1
        for ci in sorted(cnt):
            out.append({"group": g, "code": ci, "label": cb.issue.get((g, ci), f"{g}{ci}"),
                        "count": cnt[ci], "pct": _r(100 * cnt[ci] / n_resp if n_resp else 0, 1)})

    tally(a_col, "A")
    tally(b_col, "B")
    out.sort(key=lambda x: -x["pct"])
    return {"items": out, "n_respondents": n_resp}


_DISPATCH = {
    "macro": lambda df, cols, cb, q: _agg_macro(df, cols, cb),
    "mean": lambda df, cols, cb, q: _agg_mean(df, cols, cb),
    "tenure_ym": lambda df, cols, cb, q: _agg_tenure(df, cols, cb),
    "mean_pct": lambda df, cols, cb, q: _agg_mean_pct(df, cols, cb),
    "mean_pct_cf": lambda df, cols, cb, q: _agg_mean_pct_cf(df, cols, cb),
    "dist": lambda df, cols, cb, q: _agg_dist(df, cols, cb, q),
    "dist_multi": lambda df, cols, cb, q: _agg_dist_multi(df, cols, cb),
    "multi_rate": lambda df, cols, cb, q: _agg_multi_rate(df, cols, cb),
    "likert": lambda df, cols, cb, q: _agg_likert(df, cols, cb),
    "likert_iv": lambda df, cols, cb, q: _agg_likert_iv(df, cols, cb),
    "issue_rank": lambda df, cols, cb, q: _agg_issue_rank(df, q, cb),
    "issue_multi": lambda df, cols, cb, q: _agg_issue_multi(df, q, cb),
}


# ─────────────────────────────────────────────
# 공개 API
# ─────────────────────────────────────────────

def load_questions(path: Path = CONFIG) -> list[dict]:
    return yaml.safe_load(path.read_text(encoding="utf-8"))["questions"]


# ─────────────────────────────────────────────
# 업종 세그먼트 프레임워크 (segments.yaml industry — 재분류 후 자동 활성)
# ─────────────────────────────────────────────

def _industry_labelcodes(seg_yaml: dict) -> list[tuple[str, list[str]]]:
    """업종 라벨 → R_Cate 코드 리스트 (options_order 순)."""
    cfg = seg_yaml.get("industry") or {}
    gid = cfg.get("column_group", 1)
    grp = next((g for g in seg_yaml.get("column_groups", []) if g.get("id") == gid), {})
    opts = grp.get("options", {})            # code(str) → label
    order = grp.get("options_order", list(dict.fromkeys(opts.values())))
    out = []
    for lab in order:
        codes = [str(k) for k, v in opts.items() if v == lab]
        if codes:
            out.append((lab, codes))
    return out


def industry_availability(corp_df: pd.DataFrame, seg_yaml: dict) -> dict:
    """업종별 N 현황(재분류 게이트 placeholder 용). 항상 계산."""
    cfg = seg_yaml.get("industry") or {}
    src = cfg.get("source", "R_Cate")
    min_n = cfg.get("min_n", 10)
    cate = corp_df[src].astype(str).str.strip()
    counts = {lab: int(cate.isin(codes).sum()) for lab, codes in _industry_labelcodes(seg_yaml)}
    return {"enabled": bool(cfg.get("enabled")), "min_n": min_n, "source": src,
            "counts": counts, "ready": [k for k, v in counts.items() if v >= min_n]}


def industry_segments(corp_df: pd.DataFrame, seg_yaml: dict) -> list[dict]:
    """활성(enabled) & N≥min_n 업종 세그먼트 [{name, codes, n}]. 비활성이면 []."""
    cfg = seg_yaml.get("industry") or {}
    if not cfg.get("enabled"):
        return []
    src = cfg.get("source", "R_Cate")
    min_n = cfg.get("min_n", 10)
    cate = corp_df[src].astype(str).str.strip()
    out = []
    for lab, codes in _industry_labelcodes(seg_yaml):
        n = int(cate.isin(codes).sum())
        if n >= min_n:
            out.append({"name": lab, "codes": codes, "n": n})
    return out


def aggregate_question(df_seg: pd.DataFrame, q: dict, cb: Codebook) -> dict:
    cols = cb.resolve_cols(q["cols"], set(df_seg.columns))
    fn = _DISPATCH.get(q["agg"])
    if fn is None:
        data = {"note": f"unknown agg: {q['agg']}"}
    elif not cols and q["agg"] not in ("issue_rank", "issue_multi"):
        data = {"note": "no_columns"}
    else:
        data = fn(df_seg, cols, cb, q)
    return {"id": q["id"], "title": q["title"], "page": q["page"], "agg": q["agg"],
            "html": q.get("html"), "kim21": q.get("kim21"), "N": int(len(df_seg)),
            "cols": cols, "data": data}


def aggregate_segment(df_seg: pd.DataFrame, questions: list[dict], cb: Codebook) -> dict:
    """세그먼트 df + 문항정의 → {qid: 집계결과}."""
    return {q["id"]: aggregate_question(df_seg, q, cb) for q in questions}


# ─────────────────────────────────────────────
# CLI — sanity (전체 기업 세그먼트 집계 + 정합 점검)
# ─────────────────────────────────────────────

def _corporate(df: pd.DataFrame) -> pd.DataFrame:
    rdiv = num(df["R_Div"])
    emp = num(df["p1)R_1_2_2"])
    budget = num(df["p1)R_1_4_2"]) if "p1)R_1_4_2" in df.columns else None
    spend = num(df["p1)R_1_5_2"]) if "p1)R_1_5_2" in df.columns else None
    m = rdiv.isin([1, 2, 3]) & emp.notna()
    if budget is not None and spend is not None:
        m &= (budget.notna() | spend.notna())
    return df[m].copy()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", type=Path, default=DEFAULT_SRC)
    ap.add_argument("--src-sheet", default="Sheet1")
    ap.add_argument("--out", type=Path, default=None, help="JSON 덤프(선택)")
    args = ap.parse_args()

    cb = Codebook()
    questions = load_questions()
    df = pd.read_excel(args.src, sheet_name=args.src_sheet, dtype=object)
    pop = _corporate(df)
    print(f"입력 {df.shape} → 기업 모집단 N={len(pop)}  · 문항 {len(questions)}개")

    res = aggregate_segment(pop, questions, cb)

    # 요약 출력
    print("\n=== 문항별 집계 요약 ===")
    for qid, r in res.items():
        d = r["data"]
        note = d.get("note")
        head = f"  {r['page']}·{qid:14s} [{r['agg']:11s}] cols={len(r['cols']):2d}"
        if note:
            print(head + f"  ⚠ {note}")
        elif r["agg"] in ("macro", "mean", "tenure_ym"):
            print(head + "  " + " / ".join(f"{x['label'][:10]}={x.get('mean', x.get('mean_months'))}"
                                           for x in d["items"][:4]))
        elif r["agg"] in ("mean_pct", "mean_pct_cf"):
            n = len(d["items"])
            extra = f" base현={d.get('base_N_cur')}/향={d.get('base_N_fut')}" if r["agg"] == "mean_pct_cf" else f" base={d.get('base_N')} sum={d.get('sum')}"
            print(head + f"  items={n}{extra}")
        elif r["agg"] == "dist_multi":
            top = max(d["items"], key=lambda x: x.get("impl_rate", 0), default=None)
            print(head + f"  items={len(d['items'])}"
                  + (f"  최고실시율={top['label'][:10]}({top['impl_rate']}%)" if top else ""))
        elif r["agg"] in ("dist", "multi_rate"):
            n = len(d["items"])
            top = max(d["items"], key=lambda x: x.get("pct", 0), default=None)
            print(head + f"  items={n}" + (f"  top={top['label'][:12]}({top['pct']}%)" if top else ""))
        elif r["agg"] == "likert":
            print(head + "  " + " / ".join(f"{x['label'][:8]}={x['mean']}" for x in d["items"][:3]))
        elif r["agg"] == "likert_iv":
            print(head + f"  mode={d.get('mode')} items={len(d['items'])}")
        elif r["agg"] == "issue_rank":
            top = d["themes"][0] if d["themes"] else None
            print(head + f"  themes={len(d['themes'])} n_resp={d['n_respondents']}"
                  + (f"  1위={top['label'][:14]}(w{top['weighted']})" if top else ""))
        elif r["agg"] == "issue_multi":
            print(head + f"  items={len(d['items'])} n_resp={d['n_respondents']}")

    # 정합 점검: Ⅰ macro 5지표를 compute_page_data p5 metrics(동일 label 순서) 와 대조
    from stage5_render_html import compute_page_data
    pd_data = compute_page_data(pop)
    print("\n=== 정합 점검 (aggregator macro vs compute_page_data p5) ===")
    ok = True
    for agg_item, cpd in zip(res["macro"]["data"]["items"], pd_data["p5"]["metrics"]):
        a = agg_item["mean"]
        b = round(float(cpd["avg"]), 2) if cpd["avg"] is not None else None
        match = (a == b)
        ok &= match
        print(f"  {cpd['label'][:14]:14s} agg={a}  cpd={b}  {'OK' if match else '✗ 불일치'}")
    print(f"  → macro 정합: {'전부 일치' if ok else '불일치 있음'}")
    print(f"  교육투자 만족도 likert mean={res['satisfaction']['data']['items'][0]['mean']}")

    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"\nJSON 덤프: {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
