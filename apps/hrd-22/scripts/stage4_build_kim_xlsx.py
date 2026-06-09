"""[단계 4] 정제 DB → 김태웅 형식 16시트 빈도·비율 xlsx 자동 생성 (v1).

v1 범위:
  - 7시트: 전체 + 인력규모 5 + 공공조직
  - 4개 컬럼 그룹: 1)업종 / 2)인력규모 / 3)예상매출 / 4)교육예산

사용:
    # 21회 검증
    python scripts/stage4_build_kim_xlsx.py \\
        --src reference/21th_cleaned.xlsx \\
        --out workspace/test_kim_21th_v1.xlsx

    # 22회 (정제 DB 도착 후)
    python scripts/stage4_build_kim_xlsx.py \\
        --src 3_cleaned/All_확인완료_22th.xlsx \\
        --out 4_aggregated/김태웅_22th.xlsx
"""
from __future__ import annotations
import argparse
import io
import sys
from pathlib import Path

import numpy as np
import openpyxl as op
import pandas as pd
import yaml
from openpyxl.styles import Alignment, Font, PatternFill

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from segment_aggregate import Codebook, aggregate_segment, load_questions, industry_segments  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SEGMENTS_YAML = ROOT / "config/segments.yaml"
DEFAULT_SRC = ROOT / "reference/21th_cleaned.xlsx"
DEFAULT_OUT = ROOT / "workspace/test_kim_21th_v1.xlsx"

KIM_FREQ_HEADER_ROW = 4   # row 4 헤더 (label row), 5~22 빈도
KIM_FREQ_DATA_START = 5
KIM_TOTAL_ROW = 24
KIM_RATIO_HEADER_ROW = 28
KIM_RATIO_DATA_START = 29
KIM_RATIO_TOTAL_ROW = 48


# ─────────────────────────────────────────────
# 시트 필터 적용
# ─────────────────────────────────────────────

def _safe_sheet_name(name: str) -> str:
    """엑셀 시트명 제약(/ \\ ? * [ ] : 금지, 31자)."""
    for ch in r'/\?*[]:':
        name = name.replace(ch, "·")
    return name[:31]


def apply_sheet_filter(df: pd.DataFrame, sheet_def: dict, emp_col: str = "p1)R_1_2_2") -> pd.DataFrame:
    f = sheet_def["filter"]
    if f == "all":
        return df
    if f == "emp_range":
        emp = pd.to_numeric(df[emp_col], errors="coerce")
        lo = sheet_def.get("min_inclusive")
        hi = sheet_def.get("max_exclusive")
        mask = emp.notna()
        if lo is not None:
            mask &= emp >= lo
        if hi is not None:
            mask &= emp < hi
        return df[mask].copy()
    if f == "rdiv_in":
        vals = sheet_def["values"]
        rdiv = pd.to_numeric(df["R_Div"], errors="coerce")
        return df[rdiv.isin(vals)].copy()
    raise ValueError(f"unknown filter: {f}")


# ─────────────────────────────────────────────
# 컬럼 그룹별 빈도 계산
# ─────────────────────────────────────────────

def compute_group_freq(df: pd.DataFrame, group: dict) -> dict[str, int]:
    """그룹 라벨 → 빈도 dict (options_order 순)."""
    if group["type"] == "categorical":
        src = group["source"]
        order = group["options_order"]
        opts = group["options"]
        public_rule = group.get("public_rule", {})
        # 업종 처리: R_Cate 코드 → 한글 라벨, R_Cate=NaN + R_Div=4/5 → 공공
        labels = []
        for _, row in df.iterrows():
            r_cate = row.get(src)
            r_div = row.get("R_Div")
            if pd.isna(r_cate) or str(r_cate).strip() == "":
                # 공공/학교 분류
                try:
                    r_div_n = int(float(r_div))
                except (TypeError, ValueError):
                    r_div_n = None
                if r_div_n in public_rule:
                    labels.append(public_rule[r_div_n])
                    continue
                labels.append(None)
            else:
                key = str(r_cate).strip()
                lbl = opts.get(key)
                if lbl is None:
                    # 숫자 코드 normalize
                    try:
                        lbl = opts.get(str(int(float(key))))
                    except (TypeError, ValueError):
                        pass
                labels.append(lbl)
        counts = pd.Series(labels).value_counts()
        return {opt: int(counts.get(opt, 0)) for opt in order}

    if group["type"] == "numeric_range":
        src = group["source"]
        s = pd.to_numeric(df[src], errors="coerce")
        bins = group["bins"]
        labels = group["labels"]
        # null 끝값을 inf 로
        bins_use = [b if b is not None else np.inf for b in bins]
        cats = pd.cut(s, bins=bins_use, labels=labels, right=False, include_lowest=True)
        counts = cats.value_counts()
        return {lbl: int(counts.get(lbl, 0)) for lbl in labels}

    raise ValueError(f"unknown group type: {group['type']}")


# ─────────────────────────────────────────────
# 김태웅 시트에 쓰기
# ─────────────────────────────────────────────

def write_sheet(ws, freq_by_group: dict[int, dict[str, int]], groups: list[dict], total: int):
    """김태웅 시트 row 4 헤더 + row 5~ 빈도 + row 24 합계 + row 28 헤더 + row 29~ 비율 + row 48 합계."""
    # 크로스탭 열을 갖는 그룹만 (sheet_col_offset 없는 라벨 전용 그룹 0)조직구분 제외)
    groups = [g for g in groups if "sheet_col_offset" in g]
    # 헤더 (빈도) — 김태웅은 1열 라벨 + 1열 값 쌍
    ws.cell(row=KIM_FREQ_HEADER_ROW, column=1, value="구분")
    for g in groups:
        offset = g["sheet_col_offset"]
        ws.cell(row=KIM_FREQ_HEADER_ROW, column=offset + 1, value=g["label"])  # openpyxl is 1-indexed

    # 빈도
    for g in groups:
        gid = g["id"]
        offset = g["sheet_col_offset"]
        labels = g["options_order"] if g["type"] == "categorical" else g["labels"]
        freq = freq_by_group[gid]
        for i, lbl in enumerate(labels):
            r = KIM_FREQ_DATA_START + i
            ws.cell(row=r, column=offset + 1, value=lbl)
            ws.cell(row=r, column=offset + 2, value=freq.get(lbl, 0))

    # 합계 행 (빈도)
    ws.cell(row=KIM_TOTAL_ROW, column=1, value="합계")
    for g in groups:
        offset = g["sheet_col_offset"]
        ws.cell(row=KIM_TOTAL_ROW, column=offset + 2, value=total)

    # 헤더 (비율)
    ws.cell(row=KIM_RATIO_HEADER_ROW, column=1, value="구분")
    for g in groups:
        offset = g["sheet_col_offset"]
        ws.cell(row=KIM_RATIO_HEADER_ROW, column=offset + 1, value=g["label"])

    # 비율
    for g in groups:
        gid = g["id"]
        offset = g["sheet_col_offset"]
        labels = g["options_order"] if g["type"] == "categorical" else g["labels"]
        freq = freq_by_group[gid]
        for i, lbl in enumerate(labels):
            r = KIM_RATIO_DATA_START + i
            ws.cell(row=r, column=offset + 1, value=lbl)
            ratio = (freq.get(lbl, 0) / total) if total > 0 else 0
            ws.cell(row=r, column=offset + 2, value=ratio)

    # 비율 합계
    ws.cell(row=KIM_RATIO_TOTAL_ROW, column=1, value="합계")
    for g in groups:
        offset = g["sheet_col_offset"]
        ws.cell(row=KIM_RATIO_TOTAL_ROW, column=offset + 2, value=1.0 if total > 0 else 0)


# ─────────────────────────────────────────────
# 문항 집계 블록 (③ segment_aggregate 데이터 → 시트 하단 행 블록)
# ─────────────────────────────────────────────

ROMAN = {1: "Ⅰ", 2: "Ⅱ", 3: "Ⅲ", 4: "Ⅳ", 5: "Ⅴ"}
_TITLE_FILL = PatternFill("solid", fgColor="1F3A5F")
_HDR_FILL = PatternFill("solid", fgColor="E8EDF3")
_TITLE_FONT = Font(bold=True, color="FFFFFF", size=11)
_HDR_FONT = Font(bold=True, size=9)
_PCT_FMT = "0.0"


def _w(ws, r, c, v, *, font=None, fill=None, fmt=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = Alignment(horizontal=align)
    return cell


def _hdr_row(ws, r, labels):
    for i, lab in enumerate(labels):
        _w(ws, r, 1 + i, lab, font=_HDR_FONT, fill=_HDR_FILL)


def write_question_block(ws, r: int, res: dict) -> int:
    """문항 하나의 집계(res)를 row r 부터 기록, 다음 빈 행 반환."""
    agg, d = res["agg"], res["data"]
    title = f"[{ROMAN.get(res['page'], res['page'])}] {res['title']}  ·  N={res['N']}"
    _w(ws, r, 1, title, font=_TITLE_FONT, fill=_TITLE_FILL)
    for c in range(2, 9):
        ws.cell(row=r, column=c).fill = _TITLE_FILL
    r += 1

    if d.get("note"):
        _w(ws, r, 1, f"(집계 불가: {d['note']})")
        return r + 2

    if agg in ("macro", "mean"):
        unit = d.get("unit", "")
        _hdr_row(ws, r, ["구분", f"평균{('('+unit+')') if unit else ''}", "중앙값", "N"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["mean"]); _w(ws, r, 3, it["median"]); _w(ws, r, 4, it["n"]); r += 1

    elif agg == "tenure_ym":
        _hdr_row(ws, r, ["구분", "평균(개월)", "중앙(개월)", "평균(년)", "N"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["mean_months"]); _w(ws, r, 3, it["median_months"])
            _w(ws, r, 4, it["mean_years"]); _w(ws, r, 5, it["n"]); r += 1

    elif agg == "mean_pct":
        _hdr_row(ws, r, ["구분", "평균 비율(%)"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["pct"], fmt=_PCT_FMT); r += 1
        _w(ws, r, 1, f"합계(base N={d.get('base_N')})"); _w(ws, r, 2, d.get("sum"), fmt=_PCT_FMT); r += 1

    elif agg == "mean_pct_cf":
        _hdr_row(ws, r, ["구분", "현재(%)", "향후(%)", "Δ(향후-현재)"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, it["label"] + (" ※중복합산" if it.get("merged") else ""))
            _w(ws, r, 2, it["cur"], fmt=_PCT_FMT); _w(ws, r, 3, it["fut"], fmt=_PCT_FMT)
            _w(ws, r, 4, it["delta"], fmt="+0.0;-0.0"); r += 1
        _w(ws, r, 1, f"base 현재 N={d.get('base_N_cur')} / 향후 N={d.get('base_N_fut')}"); r += 1

    elif agg == "dist":
        _hdr_row(ws, r, ["응답", "빈도", "비율(%)"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["count"]); _w(ws, r, 3, it["pct"], fmt=_PCT_FMT); r += 1

    elif agg == "dist_multi":
        # 제도 × 수준 교차 + 실시율
        _hdr_row(ws, r, ["제도", "미고려", "고려중", "개발중", "실시중", "중단/폐지", "실시율(%)"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, it["label"])
            dm = {x["label"]: x["pct"] for x in it["dist"]}
            for ci, lab in enumerate(["미고려", "고려중", "개발중", "실시중", "중단/폐지"]):
                _w(ws, r, 2 + ci, dm.get(lab, 0.0), fmt=_PCT_FMT)
            _w(ws, r, 7, it["impl_rate"], fmt=_PCT_FMT); r += 1

    elif agg == "multi_rate":
        _hdr_row(ws, r, ["항목", "빈도", f"선택률(%) · 응답자 {d.get('n_respondents')}"]); r += 1
        for it in sorted(d["items"], key=lambda x: -x["pct"]):
            _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["count"]); _w(ws, r, 3, it["pct"], fmt=_PCT_FMT); r += 1

    elif agg == "likert":
        _hdr_row(ws, r, ["항목", "평균", "1점%", "2점%", "3점%", "4점%", "5점%", "N"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["mean"])
            for k in range(1, 6):
                _w(ws, r, 2 + k, it["dist"].get(k, 0.0), fmt=_PCT_FMT)
            _w(ws, r, 8, it["n"]); r += 1

    elif agg == "likert_iv":
        if d.get("mode") == "pairs":
            _hdr_row(ws, r, ["항목", "중요도", "수행도", "gap(중요-수행)", "N"]); r += 1
            for it in d["items"]:
                _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["imp_mean"]); _w(ws, r, 3, it["perf_mean"])
                _w(ws, r, 4, it["gap"], fmt="+0.00;-0.00"); _w(ws, r, 5, it["n"]); r += 1
        else:
            _hdr_row(ws, r, ["항목", "평균", "1점%", "2점%", "3점%", "4점%", "5점%", "N"]); r += 1
            for it in d["items"]:
                _w(ws, r, 1, it["label"]); _w(ws, r, 2, it["mean"])
                for k in range(1, 6):
                    _w(ws, r, 2 + k, it["dist"].get(k, 0.0), fmt=_PCT_FMT)
                _w(ws, r, 8, it["n"]); r += 1
            if d.get("gap") is not None:
                _w(ws, r, 1, f"gap({d['items'][0]['label']}-{d['items'][1]['label']})")
                _w(ws, r, 2, d["gap"], fmt="+0.00;-0.00"); r += 1

    elif agg == "issue_rank":
        _hdr_row(ws, r, ["주요 이슈(그룹)", "1순위", "2순위", "3순위", "합", "언급률(%)", "가중점수"]); r += 1
        for t in d["themes"]:
            _w(ws, r, 1, f"{t['label']} ({t['group']})"); _w(ws, r, 2, t["r1"]); _w(ws, r, 3, t["r2"])
            _w(ws, r, 4, t["r3"]); _w(ws, r, 5, t["total"]); _w(ws, r, 6, t["mention_pct"], fmt=_PCT_FMT)
            _w(ws, r, 7, t["weighted"]); r += 1
        _w(ws, r, 1, f"응답자 N={d.get('n_respondents')}"); r += 1

    elif agg == "issue_multi":
        _hdr_row(ws, r, ["기타 이슈(그룹)", "빈도", f"선택률(%) · 응답자 {d.get('n_respondents')}"]); r += 1
        for it in d["items"]:
            _w(ws, r, 1, f"{it['label']} ({it['group']})"); _w(ws, r, 2, it["count"]); _w(ws, r, 3, it["pct"], fmt=_PCT_FMT); r += 1

    else:
        _w(ws, r, 1, f"(미지원 agg: {agg})"); r += 1

    return r + 1   # 블록 사이 빈 행


def write_question_blocks(ws, agg_res: dict, questions: list[dict], start_row: int) -> int:
    """config 순서대로 전 문항 블록 기록."""
    _w(ws, start_row, 1, "■ 설문 문항별 집계 (2026 전 문항)", font=Font(bold=True, size=13))
    r = start_row + 2
    for q in questions:
        res = agg_res.get(q["id"])
        if res is None:
            continue
        r = write_question_block(ws, r, res)
    return r


# ─────────────────────────────────────────────
# main
# ─────────────────────────────────────────────

def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--src", type=Path, default=DEFAULT_SRC,
                   help="정제 DB xlsx (Sheet1 또는 첫 시트)")
    p.add_argument("--src-sheet", default="Sheet1")
    p.add_argument("--out", type=Path, default=DEFAULT_OUT)
    p.add_argument("--segments", type=Path, default=SEGMENTS_YAML)
    p.add_argument("--no-questions", action="store_true",
                   help="설문 문항별 집계 블록 생략 (21회 검증 등 인구통계만 필요할 때)")
    args = p.parse_args()

    seg = yaml.safe_load(args.segments.read_text(encoding="utf-8"))
    df_all = pd.read_excel(args.src, sheet_name=args.src_sheet, dtype=object)
    print(f"정제 DB: {df_all.shape}")

    with_q = not args.no_questions
    cb = Codebook() if with_q else None
    questions = load_questions() if with_q else []

    # 모집단 필터 — 기업/공공 분리 (docs/POPULATION_FILTER_SPEC_22th.md)
    pf = seg["population_filter"]
    excl_col = pf.get("optional_exclusion_column")
    excl_val = pf.get("optional_exclusion_value")

    def build_pop(spec: dict) -> pd.DataFrame:
        m = pd.Series(True, index=df_all.index)
        if "rdiv_in" in spec:
            rdiv = pd.to_numeric(df_all["R_Div"], errors="coerce")
            m &= rdiv.isin(spec["rdiv_in"])
        for c in spec.get("required_all", []):
            m &= pd.to_numeric(df_all[c], errors="coerce").notna()
        if spec.get("required_any"):
            any_m = pd.Series(False, index=df_all.index)
            for c in spec["required_any"]:
                any_m |= pd.to_numeric(df_all[c], errors="coerce").notna()
            m &= any_m
        if excl_col and excl_col in df_all.columns:
            m &= ~(df_all[excl_col].astype(str).str.lower() == str(excl_val).lower())
        return df_all[m].copy()

    pops = {"corporate": build_pop(pf["corporate"]), "public": build_pop(pf["public"])}
    print(f"모집단 — 기업(corporate): {len(pops['corporate'])} / 공공(public): {len(pops['public'])}")

    # 시트별 집계
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb = op.Workbook()
    wb.remove(wb.active)

    summary_rows = []
    for sheet_def in seg["sheets"]:
        base = pops[sheet_def.get("population", "corporate")]
        seg_df = apply_sheet_filter(base, sheet_def)
        n = len(seg_df)

        freq_by_group = {}
        for g in seg["column_groups"]:
            freq_by_group[g["id"]] = compute_group_freq(seg_df, g)

        ws = wb.create_sheet(sheet_def["name"])
        ws.column_dimensions["A"].width = 34
        write_sheet(ws, freq_by_group, seg["column_groups"], n)
        if with_q:
            agg_res = aggregate_segment(seg_df, questions, cb)
            last = write_question_blocks(ws, agg_res, questions, KIM_RATIO_TOTAL_ROW + 3)
            print(f"  [{sheet_def['name']}] N = {n}  · 문항블록 {len(questions)}개 (~row {last})")
        else:
            print(f"  [{sheet_def['name']}] N = {n}")
        summary_rows.append({"sheet": sheet_def["name"], "N": n})

    # 업종 시트 (industry.enabled & N≥min_n) — 재분류 후 자동 추가
    corp = pops["corporate"]
    isrc = (seg.get("industry") or {}).get("source", "R_Cate")
    icate = corp[isrc].astype(str).str.strip()
    for ind in industry_segments(corp, seg):
        seg_df = corp[icate.isin([str(c) for c in ind["codes"]])].copy()
        n = len(seg_df)
        freq_by_group = {g["id"]: compute_group_freq(seg_df, g) for g in seg["column_groups"]}
        ws = wb.create_sheet(_safe_sheet_name(ind["name"]))
        ws.column_dimensions["A"].width = 34
        write_sheet(ws, freq_by_group, seg["column_groups"], n)
        if with_q:
            write_question_blocks(ws, aggregate_segment(seg_df, questions, cb), questions, KIM_RATIO_TOTAL_ROW + 3)
        summary_rows.append({"sheet": ind["name"], "N": n})
        print(f"  [업종:{ind['name']}] N = {n}")

    wb.save(args.out)
    try:
        shown = args.out.resolve().relative_to(ROOT)
    except ValueError:
        shown = args.out
    print(f"\n출력: {shown}")

    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
