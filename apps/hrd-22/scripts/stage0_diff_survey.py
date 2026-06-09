"""[단계 0] 22회 설문 초안 텍스트 추출 + 21회 baseline 과 변경분 비교.

처리:
  1) 0_survey/survey_draft_v0.2.pptx → 0_survey/survey_22th_extracted.txt 추출
  2) reference/21th_survey_extracted.txt (21회 baseline) 과 slide 단위 diff
  3) 0_survey/change_log_21_22.md 리포트 작성

산출:
  - 0_survey/survey_22th_extracted.txt
  - 0_survey/change_log_21_22.md

사용:
    python scripts/stage0_diff_survey.py
    python scripts/stage0_diff_survey.py --pptx 0_survey/survey_draft_v0.1.pptx  # 다른 버전
"""
from __future__ import annotations

import argparse
import difflib
import io
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    sys.exit("python-pptx 미설치 — pip install python-pptx")

# Windows cp949 회피
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PPTX = ROOT / "0_survey/survey_draft_v0.2.pptx"
EXTRACTED_OUT = ROOT / "0_survey/survey_22th_extracted.txt"
BASELINE_TXT = ROOT / "reference/21th_survey_extracted.txt"
DIFF_OUT = ROOT / "0_survey/change_log_21_22.md"

SLIDE_HDR = re.compile(r"^## slide (\d+)\s*$")


# ---------- 추출 ----------

def extract_pptx_text(pptx_path: Path, out_path: Path) -> int:
    """PPTX 의 텍스트·표를 slide 단위로 추출."""
    prs = Presentation(str(pptx_path))
    lines: list[str] = []
    lines.append(f"# {pptx_path.name}")
    lines.append(f"# total slides: {len(prs.slides)}")
    lines.append("")
    for idx, slide in enumerate(prs.slides, 1):
        slide_lines = [f"\n## slide {idx}"]
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    text = "".join(r.text for r in p.runs).strip()
                    if text:
                        slide_lines.append(f"  {text}")
            if sh.has_table:
                first = sh.table.rows[0].cells[0].text[:30] if sh.table.rows else ""
                slide_lines.append(f"  [TABLE {first}]")
                for row in sh.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    slide_lines.append(f"    | {' | '.join(cells)} |")
        if len(slide_lines) > 1:
            lines.extend(slide_lines)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return len(lines)


# ---------- diff ----------

def parse_slides(text: str) -> dict[int, list[str]]:
    out: dict[int, list[str]] = {}
    cur_id: int | None = None
    cur: list[str] = []
    for line in text.splitlines():
        m = SLIDE_HDR.match(line)
        if m:
            if cur_id is not None:
                out[cur_id] = cur
            cur_id = int(m.group(1))
            cur = []
        else:
            if cur_id is not None:
                cur.append(line)
    if cur_id is not None:
        out[cur_id] = cur
    return out


def render_diff(base: dict[int, list[str]], new: dict[int, list[str]],
                base_label: str, new_label: str) -> tuple[str, dict]:
    base_ids = set(base.keys())
    new_ids = set(new.keys())

    only_new = sorted(new_ids - base_ids)
    only_base = sorted(base_ids - new_ids)
    common = sorted(base_ids & new_ids)

    changed: list[tuple[int, str]] = []
    same = 0
    for sid in common:
        b_lines = base[sid]
        n_lines = new[sid]
        if b_lines == n_lines:
            same += 1
            continue
        diff = "\n".join(difflib.unified_diff(
            b_lines, n_lines,
            fromfile=f"{base_label} s{sid}",
            tofile=f"{new_label} s{sid}", n=5,
        ))
        changed.append((sid, diff))

    summary = {
        "base_slides": len(base_ids),
        "new_slides": len(new_ids),
        "same": same,
        "changed": len(changed),
        "new_only": len(only_new),
        "base_only": len(only_base),
    }

    lines: list[str] = []
    lines.append(f"# 22회 vs 21회 설문지 변경분")
    lines.append("")
    lines.append(f"> 자동 생성: `scripts/stage0_diff_survey.py`")
    lines.append(f"> base: `{base_label}` / new: `{new_label}`")
    lines.append("")
    lines.append("## 요약")
    lines.append("")
    lines.append(f"| 구분 | 슬라이드 수 |")
    lines.append(f"|---|---:|")
    lines.append(f"| 21회 (base) | {summary['base_slides']} |")
    lines.append(f"| 22회 (new) | {summary['new_slides']} |")
    lines.append(f"| 동일 | {summary['same']} |")
    lines.append(f"| **변경** | **{summary['changed']}** |")
    lines.append(f"| **신규 (22회에만)** | **{summary['new_only']}** |")
    lines.append(f"| **삭제 (21회에만)** | **{summary['base_only']}** |")
    lines.append("")

    lines.append("## 분류 의사결정 영역")
    lines.append("")
    lines.append("아래 변경/신규/삭제를 검토 후 다음 5가지로 분류 (수기):")
    lines.append("")
    lines.append("- **A. 옵션 라벨 변경** — codebook `response_format` 의 옵션 텍스트 갱신")
    lines.append("- **B. sub_item 변경** — codebook `sub_item` 갱신 (예: 새 교육제도 추가)")
    lines.append("- **C. 신규 질문** — codebook 행 추가 + binding 추가")
    lines.append("- **D. 질문 삭제** — codebook 행 제거 + binding frozen 처리")
    lines.append("- **E. 옵션 수 변경** — codebook 옵션 추가/제거 + 차트 카테고리 영향")
    lines.append("")

    if only_new:
        lines.append(f"## 신규 슬라이드 ({len(only_new)})")
        lines.append("")
        for sid in only_new:
            lines.append(f"### slide {sid}")
            lines.append("```")
            lines.extend(new[sid])
            lines.append("```")
            lines.append("")

    if only_base:
        lines.append(f"## 삭제 슬라이드 ({len(only_base)})")
        lines.append("")
        for sid in only_base:
            lines.append(f"### slide {sid}")
            lines.append("```")
            lines.extend(base[sid])
            lines.append("```")
            lines.append("")

    if changed:
        lines.append(f"## 변경 슬라이드 ({len(changed)})")
        lines.append("")
        for sid, diff in changed:
            lines.append(f"### slide {sid}")
            lines.append("```diff")
            lines.append(diff)
            lines.append("```")
            lines.append("")

    return "\n".join(lines) + "\n", summary


# ---------- main ----------

def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="22회 설문 초안 추출 + 21회 변경분 diff")
    p.add_argument("--pptx", type=Path, default=DEFAULT_PPTX,
                   help=f"22회 설문 PPTX (기본: {DEFAULT_PPTX.name})")
    p.add_argument("--baseline", type=Path, default=BASELINE_TXT,
                   help=f"21회 baseline 텍스트 (기본: reference/21th_survey_extracted.txt)")
    p.add_argument("--extracted-out", type=Path, default=EXTRACTED_OUT)
    p.add_argument("--diff-out", type=Path, default=DIFF_OUT)
    args = p.parse_args(argv)

    if not args.pptx.is_file():
        print(f"[error] PPTX 없음: {args.pptx}", file=sys.stderr)
        return 2
    if not args.baseline.is_file():
        print(f"[error] baseline 텍스트 없음: {args.baseline}", file=sys.stderr)
        return 2

    # 1) 추출
    n_lines = extract_pptx_text(args.pptx, args.extracted_out)
    print(f"[1/2] 추출: {args.pptx.name} → {args.extracted_out.relative_to(ROOT)} ({n_lines} lines)")

    # 2) diff
    base_text = args.baseline.read_text(encoding="utf-8")
    new_text = args.extracted_out.read_text(encoding="utf-8")
    base = parse_slides(base_text)
    new = parse_slides(new_text)
    diff_md, summary = render_diff(
        base, new,
        base_label=args.baseline.name,
        new_label=args.pptx.name,
    )
    args.diff_out.parent.mkdir(parents=True, exist_ok=True)
    args.diff_out.write_text(diff_md, encoding="utf-8")
    print(f"[2/2] diff: {args.diff_out.relative_to(ROOT)}")

    print()
    print("=== 변경분 요약 ===")
    print(f"  base(21회) {summary['base_slides']} / new(22회) {summary['new_slides']} 슬라이드")
    print(f"  동일 {summary['same']} / 변경 {summary['changed']} / 신규 {summary['new_only']} / 삭제 {summary['base_only']}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
